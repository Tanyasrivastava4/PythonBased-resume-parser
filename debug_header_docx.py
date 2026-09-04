"""
debug_header_docx.py — Trace where header text disappears in DOCX extraction.

Usage:
    python debug_header_docx.py "resumes/Aman Chauhan_Machine Learning Engineer_Intileo.docx"
"""

import sys
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

_W_NS  = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_NS    = {"w": _W_NS, "mc": _MC_NS, "wp": _WP_NS}


def _short_tag(tag):
    known = {
        _W_NS:  "w",
        _MC_NS: "mc",
        _WP_NS: "wp",
        "http://schemas.openxmlformats.org/drawingml/2006/main": "a",
        "http://schemas.microsoft.com/office/word/2010/wordprocessingShape": "wps",
        "http://schemas.openxmlformats.org/drawingml/2006/picture": "pic",
        "urn:schemas-microsoft-com:vml": "v",
        "urn:schemas-microsoft-com:office:office": "o",
    }
    if tag and tag.startswith("{"):
        ns, local = tag[1:].split("}", 1)
        prefix = known.get(ns, ns[-20:])
        return f"{prefix}:{local}"
    return tag


def debug_docx(docx_path: str):
    print(f"\n{'═' * 70}")
    print(f"DEBUGGING DOCX: {Path(docx_path).name}")
    print(f"{'═' * 70}")

    doc = Document(docx_path)

    # ── 1. Document headers (<w:hdr>) ──
    print(f"\n[1] DOCUMENT PAGE-HEADERS (<w:hdr>)")
    try:
        for sec_idx, section in enumerate(doc.sections):
            hdr = section.header
            paras = hdr.paragraphs
            tables = hdr.tables
            has_text = any(p.text.strip() for p in paras)
            print(f"    Section {sec_idx}: {len(paras)} paragraphs, "
                  f"{len(tables)} tables, has_text={has_text}")
            for p_idx, para in enumerate(paras):
                t = para.text.strip()
                if t:
                    print(f"      para[{p_idx}] text: '{t[:80]}'")
                txbs = para._element.findall(".//w:txbxContent", _NS)
                if txbs:
                    print(f"      para[{p_idx}] has {len(txbs)} text box(es)")
                    for tb in txbs:
                        tb_paras = tb.findall(".//w:p", _NS)
                        tb_text = "\n".join(
                            "".join(t.text or "" for t in p.findall(".//w:t", _NS))
                            for p in tb_paras
                        ).strip()
                        print(f"        textbox content: '{tb_text[:100]}'")
    except Exception as e:
        print(f"    ERROR accessing headers: {e}")

    # ── 2. Body-level elements (first 20) ──
    body = doc.element.body
    children = list(body)
    print(f"\n[2] BODY-LEVEL ELEMENTS: {len(children)} total (showing first 20)")
    for i, el in enumerate(children[:20]):
        tag = _short_tag(el.tag)
        if el.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            text = Paragraph(el, doc).text.strip()[:60]
            drawings = el.findall(".//mc:AlternateContent", _NS)
            picts = el.findall(".//" + qn("w:pict"))
            txbs = el.findall(".//w:txbxContent", _NS)
            extras = []
            if drawings: extras.append(f"{len(drawings)} mc:AltContent")
            if picts:    extras.append(f"{len(picts)} w:pict")
            if txbs:     extras.append(f"{len(txbs)} txbxContent")
            extra_str = f"  [{', '.join(extras)}]" if extras else ""
            print(f"    [{i:2d}] {tag}  text='{text}'{extra_str}")
        elif el.tag == qn("w:tbl"):
            rows = el.findall(".//w:tr", _NS)
            print(f"    [{i:2d}] {tag}  ({len(rows)} rows)")
        elif el.tag == qn("w:sdt"):
            sdt_content = el.find(qn("w:sdtContent"))
            sdt_children = list(sdt_content) if sdt_content is not None else []
            sdt_tags = [_short_tag(c.tag) for c in sdt_children[:5]]
            print(f"    [{i:2d}] {tag}  sdtContent has {len(sdt_children)} children: {sdt_tags}")
        else:
            print(f"    [{i:2d}] {tag}")

    # ── 3. Text boxes in first 15 body paragraphs ──
    print(f"\n[3] TEXT BOXES IN BODY PARAGRAPHS (first 15)")
    body_paras = body.findall(qn("w:p"))
    for i, p_el in enumerate(body_paras[:15]):
        choices = p_el.findall(".//mc:Choice", _NS)
        all_txb = p_el.findall(".//w:txbxContent", _NS)
        fallback_txb = []
        for fb in p_el.findall(".//mc:Fallback", _NS):
            fallback_txb.extend(fb.findall(".//w:txbxContent", _NS))

        from docx.text.paragraph import Paragraph
        para_text = Paragraph(p_el, doc).text.strip()[:60]

        if all_txb or para_text:
            print(f"    para[{i}] Paragraph.text='{para_text}'")
            print(f"            mc:Choice count={len(choices)}, "
                  f"ALL txbxContent={len(all_txb)}, "
                  f"Fallback txbxContent={len(fallback_txb)}")
            for j, tb in enumerate(all_txb):
                tb_paras = tb.findall(".//w:p", _NS)
                tb_text = "\n".join(
                    "".join(t.text or "" for t in p.findall(".//w:t", _NS))
                    for p in tb_paras
                ).strip()
                in_fb = tb in fallback_txb
                parent_tags = []
                node = tb.getparent()
                depth = 0
                while node is not None and depth < 6:
                    parent_tags.append(_short_tag(node.tag))
                    node = node.getparent()
                    depth += 1
                print(f"            txb[{j}] in_fallback={in_fb} text='{tb_text[:80]}'")
                print(f"                  parents: {' > '.join(parent_tags)}")

    # ── 4. Structured document tags ──
    sdt_elements = body.findall(qn("w:sdt"))
    print(f"\n[4] STRUCTURED DOCUMENT TAGS (<w:sdt>): {len(sdt_elements)}")
    for i, sdt in enumerate(sdt_elements[:5]):
        sdt_content = sdt.find(qn("w:sdtContent"))
        if sdt_content is not None:
            for j, child in enumerate(sdt_content):
                tag = _short_tag(child.tag)
                if child.tag == qn("w:p"):
                    from docx.text.paragraph import Paragraph
                    text = Paragraph(child, doc).text.strip()[:60]
                    print(f"    sdt[{i}] child[{j}] {tag} text='{text}'")
                else:
                    print(f"    sdt[{i}] child[{j}] {tag}")

    # ── 5. Final output ──
    print(f"\n[5] FINAL read_docx() OUTPUT (first 500 chars)")
    from ingestion.docx_reader import read_docx
    try:
        final_text = read_docx(docx_path)
        print(f"{'─' * 50}")
        print(final_text[:500])
        print(f"{'─' * 50}")
        print(f"    Total length: {len(final_text)}")
        for keyword in ["AMAN", "Aman", "Chauhan", "chauhan", "6387", "gmail"]:
            if keyword.lower() in final_text.lower():
                print(f"    ✅ Found '{keyword}' in output")
            else:
                print(f"    ❌ Missing '{keyword}' from output")
    except Exception as e:
        print(f"    ERROR: {e}")

    print(f"\n{'═' * 70}\n")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python debug_header_docx.py \"resumes/Your Resume.docx\"")
        sys.exit(1)
    debug_docx(sys.argv[1])